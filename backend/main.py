"""
Folio, servicio OCR: recibe imagen o PDF (escaneado o con texto) y devuelve el texto por página.
Exporta a .docx a partir del texto (posiblemente editado por el usuario).
Además, herramientas clásicas de PDF: unir, separar, ordenar/rotar, comprimir, imágenes ↔ PDF,
marca de agua, numeración de páginas y protección con contraseña.
"""
import io
import math
import os
import re
import subprocess
import tempfile
import zipfile
from urllib.parse import quote

import pytesseract
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from pdf2image import convert_from_bytes, pdfinfo_from_bytes
from PIL import Image, ImageOps
from pydantic import BaseModel
from pypdf import PdfReader, PdfWriter, Transformation
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

MAX_MB = int(os.getenv("MAX_MB", "20"))
MAX_PAGINAS = int(os.getenv("MAX_PAGINAS", "50"))
MAX_ARCHIVOS = int(os.getenv("MAX_ARCHIVOS", "20"))
DPI = int(os.getenv("OCR_DPI", "300"))
ORIGENES = [o.strip() for o in os.getenv("CORS_ORIGINS", "http://localhost:3000").split(",")]

TIPOS_IMAGEN = {"image/png", "image/jpeg", "image/webp", "image/tiff", "image/bmp"}

app = FastAPI(title="Folio API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=ORIGENES,
    allow_methods=["POST", "GET"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],  # el frontend toma de ahí el nombre del archivo
)


def preparar(img: Image.Image) -> Image.Image:
    """Endereza según EXIF, pasa a escala de grises y sube contraste."""
    img = ImageOps.exif_transpose(img)
    img = ImageOps.grayscale(img)
    return ImageOps.autocontrast(img)


def ocr_imagen(img: Image.Image, idioma: str) -> str:
    # --psm 3: segmentación automática de página (lo normal para documentos)
    return pytesseract.image_to_string(preparar(img), lang=idioma, config="--psm 3").strip()


def texto_nativo_pdf(datos: bytes) -> list[str] | None:
    """Si el PDF ya trae capa de texto (no escaneado), úsala: es exacta y mucho más rápida."""
    try:
        lector = PdfReader(io.BytesIO(datos))
        paginas = [(p.extract_text() or "").strip() for p in lector.pages]
    except Exception:
        return None
    # Heurística: si casi todas las páginas tienen texto, se considera PDF digital
    con_texto = sum(1 for t in paginas if len(t) > 30)
    return paginas if paginas and con_texto >= len(paginas) * 0.8 else None


@app.get("/salud")
def salud():
    return {"ok": True, "idiomas": pytesseract.get_languages()}


@app.post("/ocr")
async def ocr(
    archivo: UploadFile = File(...),
    idioma: str = Form("spa"),          # "spa", "eng" o "spa+eng"
    forzar_ocr: bool = Form(False),     # ignora la capa de texto del PDF
):
    datos = await archivo.read()
    if len(datos) > MAX_MB * 1024 * 1024:
        raise HTTPException(413, f"El archivo pasa de {MAX_MB} MB.")

    tipo = archivo.content_type or ""
    nombre = (archivo.filename or "documento").rsplit(".", 1)[0]

    if tipo == "application/pdf":
        if not forzar_ocr and (nativo := texto_nativo_pdf(datos)):
            return {"nombre": nombre, "metodo": "texto-nativo", "paginas": nativo}

        total = pdfinfo_from_bytes(datos)["Pages"]
        if total > MAX_PAGINAS:
            raise HTTPException(413, f"El PDF tiene {total} páginas; el máximo es {MAX_PAGINAS}.")

        paginas = []
        # Página por página para no cargar todo el PDF rasterizado en memoria
        for n in range(1, total + 1):
            img = convert_from_bytes(datos, dpi=DPI, first_page=n, last_page=n)[0]
            paginas.append(ocr_imagen(img, idioma))
        return {"nombre": nombre, "metodo": "ocr", "paginas": paginas}

    if tipo in TIPOS_IMAGEN:
        try:
            img = Image.open(io.BytesIO(datos))
        except Exception:
            raise HTTPException(400, "No se pudo abrir la imagen.")
        # TIFF puede traer varias páginas
        paginas = []
        for i in range(getattr(img, "n_frames", 1)):
            img.seek(i)
            paginas.append(ocr_imagen(img.copy(), idioma))
        return {"nombre": nombre, "metodo": "ocr", "paginas": paginas}

    raise HTTPException(415, "Formato no soportado. Sube PDF, PNG, JPG, WEBP, TIFF o BMP.")


class SolicitudDocx(BaseModel):
    nombre: str = "documento"
    paginas: list[str]


@app.post("/docx")
def exportar_docx(body: SolicitudDocx):
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    for i, texto in enumerate(body.paginas):
        for parrafo in texto.split("\n\n"):
            # Une los saltos de línea internos del OCR para que el párrafo fluya
            doc.add_paragraph(" ".join(l.strip() for l in parrafo.splitlines() if l.strip()))
        if i < len(body.paginas) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    buf = io.BytesIO()
    doc.save(buf)
    return archivo_respuesta(
        buf.getvalue(), f"{body.nombre}.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


# ---------------------------------------------------------------------------
# Herramientas de PDF
# ---------------------------------------------------------------------------

def archivo_respuesta(datos: bytes, nombre: str, tipo: str) -> Response:
    # filename* (RFC 5987) para que nombres con tildes o ñ no rompan la cabecera
    return Response(datos, media_type=tipo,
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(nombre)}"})


def pdf_respuesta(escritor: PdfWriter, nombre: str) -> Response:
    buf = io.BytesIO()
    escritor.write(buf)
    return archivo_respuesta(buf.getvalue(), f"{nombre}.pdf", "application/pdf")


def zip_respuesta(archivos: list[tuple[str, bytes]], nombre: str) -> Response:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre_archivo, datos in archivos:
            z.writestr(nombre_archivo, datos)
    return archivo_respuesta(buf.getvalue(), f"{nombre}.zip", "application/zip")


def base(archivo: UploadFile) -> str:
    return (archivo.filename or "documento").rsplit(".", 1)[0]


def leer(archivo: UploadFile) -> bytes:
    datos = archivo.file.read()
    if len(datos) > MAX_MB * 1024 * 1024:
        raise HTTPException(413, f"«{archivo.filename}» pasa de {MAX_MB} MB.")
    return datos


def abrir_pdf(datos: bytes, nombre: str = "El archivo") -> PdfReader:
    try:
        lector = PdfReader(io.BytesIO(datos))
        # Los PDF con solo contraseña de propietario se abren con la contraseña vacía
        if lector.is_encrypted and not lector.decrypt(""):
            raise HTTPException(400, f"{nombre} está protegido con contraseña.")
        if not lector.pages:
            raise HTTPException(400, f"{nombre} no tiene páginas.")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, f"{nombre} no es un PDF válido o está dañado.")
    return lector


def interpretar_paginas(spec: str, total: int) -> list[list[int]]:
    """'1-3, 5, 8-' → [[0, 1, 2], [4], [7, …, total-1]] (base 0). Un rango invertido (5-1) va en orden descendente."""
    grupos = []
    for parte in re.split(r"[,;]", spec):
        parte = parte.replace(" ", "")
        if not parte:
            continue
        m = re.fullmatch(r"(\d+)(?:-(\d*))?", parte)
        if not m:
            raise HTTPException(400, f"No entiendo «{parte}». Usa números y rangos, por ejemplo: 1-3, 5, 8-")
        a = int(m[1])
        b = a if m[2] is None else int(m[2] or total)
        for n in (a, b):
            if not 1 <= n <= total:
                raise HTTPException(400, f"La página {n} no existe; el PDF tiene {total}.")
        paso = 1 if b >= a else -1
        grupos.append([n - 1 for n in range(a, b + paso, paso)])
    return grupos


def superponer(lector: PdfReader, dibujar) -> PdfWriter:
    """Dibuja con reportlab una capa por página (dibujar(c, ancho, alto, i)) y la estampa encima."""
    escritor = PdfWriter(clone_from=lector)
    capa = io.BytesIO()
    c = canvas.Canvas(capa)
    for i, pagina in enumerate(escritor.pages):
        # Así la capa queda derecha aunque la página tenga /Rotate
        pagina.transfer_rotation_to_content()
        caja = pagina.cropbox
        c.setPageSize((float(caja.width), float(caja.height)))
        dibujar(c, float(caja.width), float(caja.height), i)
        c.showPage()
    c.save()
    for pagina, encima in zip(escritor.pages, PdfReader(capa).pages):
        caja = pagina.cropbox
        pagina.merge_transformed_page(encima, Transformation().translate(float(caja.left), float(caja.bottom)))
    return escritor


@app.post("/pdf/unir")
def unir(archivos: list[UploadFile] = File(...)):
    if not 2 <= len(archivos) <= MAX_ARCHIVOS:
        raise HTTPException(400, f"Sube entre 2 y {MAX_ARCHIVOS} PDF.")
    escritor = PdfWriter()
    for a in archivos:
        escritor.append(abrir_pdf(leer(a), f"«{a.filename}»"))
    return pdf_respuesta(escritor, f"{base(archivos[0])}-unido")


@app.post("/pdf/organizar")
def organizar(
    archivo: UploadFile = File(...),
    paginas: str = Form(""),          # selección y orden; vacío = todas en su orden
    rotar: int = Form(0),             # 0, 90, 180 o 270 grados en sentido horario
    paginas_rotar: str = Form(""),    # números originales; vacío = todas
):
    if rotar % 90:
        raise HTTPException(400, "La rotación debe ser múltiplo de 90°.")
    lector = abrir_pdf(leer(archivo))
    total = len(lector.pages)

    orden = [i for g in interpretar_paginas(paginas, total) for i in g] or list(range(total))
    repetidas = {i + 1 for i in orden if orden.count(i) > 1}
    if repetidas:
        raise HTTPException(400, f"Cada página puede aparecer una sola vez; se repite: {', '.join(map(str, sorted(repetidas)))}.")
    a_rotar = {i for g in interpretar_paginas(paginas_rotar, total) for i in g} or set(range(total))

    escritor = PdfWriter()
    for i in orden:
        pagina = escritor.add_page(lector.pages[i])
        if rotar and i in a_rotar:
            pagina.rotate(rotar)
    return pdf_respuesta(escritor, f"{base(archivo)}-organizado")


@app.post("/pdf/separar")
def separar(
    archivo: UploadFile = File(...),
    rangos: str = Form(""),           # cada rango es un archivo; vacío = un archivo por página
):
    lector = abrir_pdf(leer(archivo))
    total = len(lector.pages)
    grupos = interpretar_paginas(rangos, total) or [[i] for i in range(total)]
    nombre = base(archivo)

    partes = []
    for g in grupos:
        escritor = PdfWriter()
        for i in g:
            escritor.add_page(lector.pages[i])
        buf = io.BytesIO()
        escritor.write(buf)
        sufijo = f"p{g[0] + 1}" if len(g) == 1 else f"p{g[0] + 1}-{g[-1] + 1}"
        partes.append((f"{nombre}-{sufijo}.pdf", buf.getvalue()))

    if len(partes) == 1:
        return archivo_respuesta(partes[0][1], partes[0][0], "application/pdf")
    return zip_respuesta(partes, f"{nombre}-separado")


# Resolución (dpi) a la que Ghostscript reduce las imágenes en cada nivel
NIVELES = {"fuerte": 100, "recomendada": 150, "ligera": 220}


@app.post("/pdf/comprimir")
def comprimir(archivo: UploadFile = File(...), nivel: str = Form("recomendada")):
    if nivel not in NIVELES:
        raise HTTPException(400, "Nivel no válido.")
    datos = leer(archivo)
    abrir_pdf(datos)
    dpi = NIVELES[nivel]

    with tempfile.TemporaryDirectory() as tmp:
        entrada, salida = os.path.join(tmp, "entrada.pdf"), os.path.join(tmp, "salida.pdf")
        with open(entrada, "wb") as f:
            f.write(datos)
        try:
            subprocess.run(
                ["gs", "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.5", "-dPDFSETTINGS=/ebook",
                 "-dNOPAUSE", "-dBATCH", "-dQUIET", "-dSAFER",
                 f"-dColorImageResolution={dpi}", f"-dGrayImageResolution={dpi}",
                 f"-dMonoImageResolution={dpi * 2}",
                 f"-sOutputFile={salida}", entrada],
                check=True, capture_output=True, timeout=300,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
            raise HTTPException(500, "No se pudo comprimir este PDF.")
        with open(salida, "rb") as f:
            comprimido = f.read()

    # Si ya estaba optimizado, devolver el original en vez de uno más pesado
    return archivo_respuesta(min(comprimido, datos, key=len), f"{base(archivo)}-comprimido.pdf", "application/pdf")


# Lado mayor máximo (px) de cada foto en el PDF: de sobra para leer e imprimir, sin pesar como el original
LADO_MAX = 2400


@app.post("/pdf/desde-imagenes")
def desde_imagenes(archivos: list[UploadFile] = File(...)):
    if not 1 <= len(archivos) <= MAX_ARCHIVOS:
        raise HTTPException(400, f"Sube entre 1 y {MAX_ARCHIVOS} imágenes.")
    escritor = PdfWriter()
    for a in archivos:
        try:
            img = Image.open(io.BytesIO(leer(a)))
        except Exception:
            raise HTTPException(400, f"«{a.filename}» no es una imagen válida.")
        for n in range(getattr(img, "n_frames", 1)):
            img.seek(n)
            pagina = ImageOps.exif_transpose(img.copy())
            if pagina.mode in ("RGBA", "LA", "PA") or "transparency" in pagina.info:
                pagina = pagina.convert("RGBA")
                fondo = Image.new("RGB", pagina.size, "white")
                fondo.paste(pagina, mask=pagina.getchannel("A"))
                pagina = fondo
            else:
                pagina = pagina.convert("RGB")
            pagina.thumbnail((LADO_MAX, LADO_MAX))
            # Resolución elegida para que el lado mayor mida lo que el de una hoja A4 (11,69 in)
            buf = io.BytesIO()
            pagina.save(buf, "PDF", resolution=max(pagina.size) / 11.69, quality=85)
            # Una página a la vez para no tener todas las fotos en memoria
            escritor.append(PdfReader(buf))
    return pdf_respuesta(escritor, base(archivos[0]) if len(archivos) == 1 else "imagenes")


@app.post("/pdf/a-imagenes")
def a_imagenes(
    archivo: UploadFile = File(...),
    formato: str = Form("jpg"),       # "jpg" o "png"
    dpi: int = Form(150),
):
    if formato not in ("jpg", "png"):
        raise HTTPException(400, "Formato no válido.")
    datos = leer(archivo)
    total = len(abrir_pdf(datos).pages)
    if total > MAX_PAGINAS:
        raise HTTPException(413, f"El PDF tiene {total} páginas; el máximo es {MAX_PAGINAS}.")
    dpi = max(72, min(dpi, 300))
    nombre = base(archivo)

    imagenes = []
    for n in range(1, total + 1):
        img = convert_from_bytes(datos, dpi=dpi, first_page=n, last_page=n)[0]
        buf = io.BytesIO()
        if formato == "jpg":
            img.convert("RGB").save(buf, "JPEG", quality=90)
        else:
            img.save(buf, "PNG", optimize=True)
        imagenes.append((f"{nombre}-p{n}.{formato}", buf.getvalue()))

    if total == 1:
        return archivo_respuesta(imagenes[0][1], imagenes[0][0], "image/jpeg" if formato == "jpg" else "image/png")
    return zip_respuesta(imagenes, f"{nombre}-imagenes")


@app.post("/pdf/marca-agua")
def marca_agua(archivo: UploadFile = File(...), texto: str = Form(...)):
    texto = texto.strip()
    if not texto:
        raise HTTPException(400, "Escribe el texto de la marca de agua.")
    fuente = "Helvetica-Bold"

    def dibujar(c, ancho, alto, _):
        # Texto en diagonal, del tamaño que llene ~3/4 de la diagonal de la hoja
        tam = min(0.75 * math.hypot(ancho, alto) / stringWidth(texto, fuente, 1), 0.25 * min(ancho, alto))
        c.saveState()
        c.setFillColorRGB(0.5, 0.5, 0.5)
        c.setFillAlpha(0.2)
        c.translate(ancho / 2, alto / 2)
        c.rotate(math.degrees(math.atan2(alto, ancho)))
        c.setFont(fuente, tam)
        c.drawCentredString(0, -tam * 0.35, texto)
        c.restoreState()

    return pdf_respuesta(superponer(abrir_pdf(leer(archivo)), dibujar), f"{base(archivo)}-marca")


FORMATOS_NUMERO = {"n": "{n}", "n-de-total": "{n} de {total}", "pagina-n-de-total": "Página {n} de {total}"}
POSICIONES = {"abajo-centro", "abajo-derecha", "abajo-izquierda", "arriba-derecha"}


@app.post("/pdf/numerar")
def numerar(
    archivo: UploadFile = File(...),
    posicion: str = Form("abajo-centro"),
    formato: str = Form("n"),
    inicio: int = Form(1),
):
    if posicion not in POSICIONES or formato not in FORMATOS_NUMERO:
        raise HTTPException(400, "Opciones no válidas.")
    lector = abrir_pdf(leer(archivo))
    ultimo = inicio + len(lector.pages) - 1
    margen = 28  # puntos (~1 cm)

    def dibujar(c, ancho, alto, i):
        t = FORMATOS_NUMERO[formato].format(n=inicio + i, total=ultimo)
        c.setFont("Helvetica", 10)
        c.setFillColorRGB(0.25, 0.25, 0.25)
        y = alto - margen if posicion.startswith("arriba") else margen - 8
        if posicion.endswith("centro"):
            c.drawCentredString(ancho / 2, y, t)
        elif posicion.endswith("derecha"):
            c.drawRightString(ancho - margen, y, t)
        else:
            c.drawString(margen, y, t)

    return pdf_respuesta(superponer(lector, dibujar), f"{base(archivo)}-numerado")


@app.post("/pdf/proteger")
def proteger(archivo: UploadFile = File(...), contrasena: str = Form(...)):
    if len(contrasena) < 4:
        raise HTTPException(400, "La contraseña debe tener al menos 4 caracteres.")
    escritor = PdfWriter(clone_from=abrir_pdf(leer(archivo)))
    escritor.encrypt(user_password=contrasena, algorithm="AES-256")
    return pdf_respuesta(escritor, f"{base(archivo)}-protegido")
